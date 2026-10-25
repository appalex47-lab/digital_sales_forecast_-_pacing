"""«Oportunidades» (propuesta B, etapa PRE): cálculo sobre datos de producto, filtros, detalle, paso a promoción, CSV, vacío sin datos. Uso: python3 batch_oportunidades.py <raiz>"""
import asyncio, sys, os, datetime as dt, threading, functools, http.server, socketserver
from playwright.async_api import async_playwright
ROOT = os.path.abspath(sys.argv[1])
class Q(http.server.SimpleHTTPRequestHandler):
    def log_message(self, *a): pass
socketserver.TCPServer.allow_reuse_address = True
s = socketserver.ThreadingTCPServer(('127.0.0.1', 0), functools.partial(Q, directory=ROOT)); threading.Thread(target=s.serve_forever, daemon=True).start(); port = s.server_address[1]
fails = []
def chk(n, ok, d=''):
    print(('✔ ' if ok else '✘ ') + n + (' — ' + str(d)[:380] if d and not ok else ''))
    if not ok: fails.append(n)
U = f'http://127.0.0.1:{port}/index.html'
async def boot(c):
    q = await c.new_page(); errs = []; q.on('pageerror', lambda e: errs.append(str(e)[:200]))
    await q.clock.set_fixed_time(dt.datetime(2026, 10, 2, 9, 0)); await q.goto(U + '#inicio'); await q.wait_for_timeout(1500)
    await q.evaluate("document.querySelectorAll('dialog').forEach(d=>d.remove())")
    return q, errs
async def body(q): return await q.evaluate("document.getElementById('pm-body').innerText")
async def main():
    async with async_playwright() as p:
        b = await p.chromium.launch(args=['--no-sandbox'])
        c = await b.new_context(viewport={'width': 1280, 'height': 900}, accept_downloads=True)
        q, errs = await boot(c)
        await q.evaluate("location.hash='#promociones'"); await q.wait_for_timeout(700)
        tabs = await q.evaluate("[...document.querySelectorAll('[data-action=\"pm-tab\"]')].map(b=>b.textContent.trim())")
        chk('O1 cinco pestañas: Resumen, Oportunidades, Ofertas recibidas…', len(tabs) == 5 and tabs[0].startswith('Resumen') and tabs[1].startswith('Oportunidades') and tabs[2].startswith('Ofertas'), tabs)
        await q.click('[data-action="pm-tab"][data-tab="oportunidades"]'); await q.wait_for_timeout(500)
        t = await body(q)
        chk('O2 sin datos de producto lo dice y no calcula', ('productos' in t.lower()) and 'Calcular' not in t.split('\n')[0] and not await q.query_selector('#op-t-det'), t[:300])
        r = await q.evaluate("""(async()=>{
          const gen = FP.mockProducts.generate({ from: '2026-08-01', to: '2026-09-30', skus: 80 });
          const s1 = FP.productImport.stage(gen.sales, { fileName: 'venta.csv', kind: 'sales' }); const r1 = await FP.productImport.process(s1, {});
          await FP.productStore.commit(r1.draft, { fileName: 'venta.csv' });
          const s2 = FP.productImport.stage(gen.funnel, { fileName: 'funnel.csv', kind: 'funnel' }); const r2 = await FP.productImport.process(s2, {});
          await FP.productStore.commit(r2.draft, { fileName: 'funnel.csv' });
          return FP.productStore.meta.dateMax;
        })()""")
        chk('O3 productos cargados', bool(r), r)
        # crear una promoción destino
        await q.evaluate("FP.app.state.pm.actor='Persona A'; const reg=FP.app.state.pm.reg; FP.promotions.createPromotion(reg,{name:'Demo oportunidades',createdBy:'Persona A',plannedStart:'2026-11-01',plannedEnd:'2026-11-10'},{year:2026}); FP.app.render()")
        await q.click('[data-action="pm-tab"][data-tab="ofertas"]'); await q.click('[data-action="pm-tab"][data-tab="oportunidades"]'); await q.wait_for_timeout(2500)
        n = await q.evaluate("FP.app.state.op.result ? FP.app.state.op.result.items.length : -1")
        chk('O4 calcula oportunidades al entrar', n >= 0, n)
        t = await body(q)
        chk('O5 muestra periodo y no promete efecto', 'No es un puntaje' in t, t[:300])
        if n > 0:
            rows = await q.evaluate("document.querySelectorAll('#pm-body input[data-action=\"op-chk\"]').length")
            chk('O6 hay filas con casilla', rows > 0, rows)
            first = await q.evaluate("FP.app.state.op.result.items[0].sku")
            await q.click(f'button[data-action="op-open"][data-id="{first}"]'); await q.wait_for_timeout(300)
            chk('O7 abre el detalle', await q.query_selector('#op-t-det') is not None or first in await body(q))
            await q.click('[data-action="op-close"]')
            await q.evaluate("FP.app.state.op.checked={}"); 
            await q.check(f'input[data-action="op-chk"][data-id="{first}"]'); await q.wait_for_timeout(200)
            pid = await q.evaluate("FP.app.state.pm.reg.promotions[0].promoId")
            await q.select_option('#op-b-promo', pid); await q.click('[data-action="op-to-promo"]'); await q.wait_for_timeout(400)
            sk = await q.evaluate("(FP.app.state.pm.reg.promotions[0].skus||[]).map(x=>[x.sku,x.origin])")
            chk('O8 agrega a la promoción con origen digital', [first, 'digital'] in sk, sk)
            async with q.expect_download() as d: await q.click('[data-action="op-export-csv"]')
            dl = await d.value; txt = open(await dl.path(), encoding='utf-8-sig').read()
            chk('O9 CSV con encabezado y filas', txt.startswith('SKU') and len(txt.splitlines()) > 1, txt[:120])
            # Excel de la propuesta
            await q.click('[data-action="op-export-xlsx"]'); await q.wait_for_timeout(300)
            chk('X1 sin nombre no exporta', 'nombre de la propuesta' in (await body(q)).lower() and 'Se exportaron' not in await body(q))
            await q.fill('#op-x-name', 'Propuesta Otoño 2026'); await q.dispatch_event('#op-x-name', 'change')
            async with q.expect_download() as d2: await q.click('[data-action="op-export-xlsx"]')
            dx = await d2.value; xp = '/tmp/claude-0/prop_test.xlsx'; await dx.save_as(xp)
            chk('X2 nombre del archivo', dx.suggested_filename.startswith('propuesta_propuesta_otono_2026_') and dx.suggested_filename.endswith('.xlsx'), dx.suggested_filename)
            import openpyxl
            wb = openpyxl.load_workbook(xp); ws = wb['Productos']
            hdr = [c.value for c in ws[6]]
            chk('X3 una sola hoja, encabezado en la fila 6', wb.sheetnames == ['Productos'] and hdr[0] == 'SKU' and hdr[3] == 'Subcategoría' and 'Marca' not in hdr, (wb.sheetnames, hdr[:3]))
            chk('X4 sin comprador ni precios; con laboratorio, origen, descuento y programa', not any(k in ' '.join(str(h) for h in hdr).lower() for k in ['comprador', 'precio']) and all(k in hdr for k in ['Laboratorio', 'Origen de la oferta', 'Descuento', 'Programa']), hdr)
            nvis = await q.evaluate("FP.opportunityView.visible(FP.app.state.op).length")
            nrows = ws.max_row - 6
            chk('X5 las filas coinciden con la pantalla', nrows == nvis or nrows == 1, (nrows, nvis))
            r7 = [c.value for c in ws[7]]
            chk('X6 números numéricos y porcentajes con formato', isinstance(r7[9], (int, float)) and ws.cell(7, 12).number_format == '0.0%' and ws.cell(7, 10).number_format == '"$"#,##0', (r7[9], ws.cell(7, 12).number_format))
            chk('X7 autofiltro y congelado', ws.auto_filter.ref and ws.freeze_panes == 'B7', (ws.auto_filter.ref, ws.freeze_panes))
            chk('X8 encabezado dice N de M', 'productos con oportunidad' in str(ws['A3'].value), ws['A3'].value)
            tp = (await q.evaluate("FP.app.state.op.result.items[0].types[0].id"))
            await q.select_option('#op-f-type', tp); await q.wait_for_timeout(200)
            chk('O10 filtro por tipo', await q.evaluate("FP.app.state.op.f.type") == tp)
            r = await q.evaluate("""(()=>{const op={f:{}, result:{items:[{sku:'A',types:[],category:'',evidence:{level:'low'},name:'',brand:'',offers:[{origin:'pricing'}]},{sku:'B',types:[],category:'',evidence:{level:'low'},name:'',brand:'',offers:[{origin:'commercial'},{origin:'pricing'}]},{sku:'C',types:[],category:'',evidence:{level:'low'},name:'',brand:'',offers:[]}]}};
              const v=(o)=>{op.f={origin:o}; return FP.opportunityView.visible(op).map(i=>i.sku).join('')}; return [v('pricing'),v('commercial'),v('other'),v('')]})()""")
            chk('O12 filtro por origen de la oferta', r == ['AB', 'B', '', 'ABC'] and await q.query_selector('#op-f-org') is not None, r)
            # O13–O19: ver por subcategoría / laboratorio y exportar la selección
            await q.select_option('#op-f-type', ''); await q.wait_for_timeout(200)
            nsku = await q.evaluate("FP.opportunityView.visible(FP.app.state.op, FP.app.state).length")
            await q.click('[data-action="op-view"][data-v="lab"]'); await q.wait_for_timeout(300)
            gl = await q.evaluate("FP.opportunityView.groupsOf(FP.app.state.op, FP.app.state).map(g=>[g.label,g.count,g.sortKey])")
            chk('O13 la vista por laboratorio agrupa todos los SKU visibles, de mayor a menor', sum(x[1] for x in gl) == nsku and all(gl[k][2] >= gl[k+1][2] for k in range(len(gl)-1)) and await q.query_selector('[data-action="op-gchk"]') is not None, (nsku, gl[:4]))
            await q.click('[data-action="op-view"][data-v="subcat"]'); await q.wait_for_timeout(300)
            gs = await q.evaluate("FP.opportunityView.groupsOf(FP.app.state.op, FP.app.state).map(g=>[g.label,g.count])")
            chk('O14 la vista por subcategoría también cuadra con la lista', sum(x[1] for x in gs) == nsku and 'Subcategoría' in await body(q), (nsku, gs[:4]))
            async with q.expect_download() as dl: await q.click('[data-action="op-export-csv"]')
            d = await dl.value; txt = open(await d.path(), encoding='utf-8-sig').read().strip().splitlines()
            chk('O15 sin marcar, el CSV agrupado trae todos los grupos visibles', d.suggested_filename.startswith('oportunidades_por_subcategoria') and len(txt) == len(gs) + 1 and txt[0].startswith('Subcategoría,'), (d.suggested_filename, len(txt), len(gs)))
            await q.evaluate("document.querySelector('[data-action=\"op-gchk\"]').click()"); await q.wait_for_timeout(300)
            first = gs[0]
            async with q.expect_download() as dl: await q.click('[data-action="op-export-csv"]')
            d = await dl.value; txt = open(await d.path(), encoding='utf-8-sig').read().strip().splitlines()
            chk('O16 marcando un grupo, el CSV trae solo ese grupo', len(txt) == 2 and (first[0] in txt[1] or first[0] == 'Sin subcategoría'), (txt[:2], first))
            await q.fill('#op-x-name', 'Prueba agrupada')
            async with q.expect_download() as dl: await q.click('[data-action="op-export-xlsx"]')
            d = await dl.value
            import openpyxl
            import tempfile; xp = os.path.join(tempfile.mkdtemp(), 'agrupado.xlsx'); await d.save_as(xp)
            wb = openpyxl.load_workbook(xp)
            chk('O17 el Excel trae los SKU del grupo marcado y la hoja «Por subcategoría»', wb.sheetnames == ['Productos', 'Por subcategoría'] and wb['Productos'].max_row - 6 == first[1] and wb['Por subcategoría'].max_row == 4 + 1, (wb.sheetnames, wb['Productos'].max_row - 6, first))
            await q.click('[data-action="op-gopen"]'); await q.wait_for_timeout(300)
            chk('O18 «Ver SKU» pasa a la lista por SKU filtrada por ese grupo', await q.evaluate("FP.app.state.op.view") == 'sku' and await q.evaluate("FP.opportunityView.visible(FP.app.state.op, FP.app.state).length") == first[1] and await q.query_selector('#op-f-sub') is not None, first)
            await q.evaluate("FP.app.state.op.f.subcategory=''; FP.app.state.op.f.lab=''; FP.app.state.op.checkedG={}; FP.app.render()"); await q.wait_for_timeout(200)
            chk('O19 sin errores de página al agrupar y exportar', not errs, errs)
        chk('O11 sin errores de página', not errs, errs)
        await b.close()
asyncio.run(main())
print('FALLAS:', fails); sys.exit(1 if fails else 0)
